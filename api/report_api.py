import io
import os
from flask import Blueprint, request, send_file
from models import Ucastie, Meropriyatie, Uroven

report_bp = Blueprint('report_bp', __name__, url_prefix='/api/report')

TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    'templates',
    'report_template.docx',
)

PED_TITLE = (
    'Участие педагогических работников в конференциях, семинарах, конкурсах '
    'профессионального мастерства регионального, федерального уровней в {period} уч году'
)
STUD_TITLE = (
    'Участие студентов в конференциях, олимпиадах, соревнованиях, конкурсах '
    'профессионального мастерства регионального, всероссийского, международного '
    'уровней в {period} уч. году'
)
TABLE_HEADERS = [
    '№',
    'Наименование мероприятия',
    'Уровень мероприятия',
    'Сроки проведения',
    'Результат ФИО участника Дипломы, награды',
]


def _is_student(u):
    return bool(u.user.group and str(u.user.group).strip())


def _parse_period(period):
    year1, year2 = period.split('-')
    return int(year1), int(year2)


def _format_result_column(u, is_stud):
    parts = [p.strip() for p in (u.rezultat or '').split(',')]
    result = parts[0] if parts else ''
    diplomas = parts[1] if len(parts) > 1 else ''
    awards = parts[2] if len(parts) > 2 else ''

    chunks = []
    if is_stud and u.user.group:
        chunks.append(f'{u.user.fio} ({u.user.group})')
    else:
        chunks.append(u.user.fio)

    if result:
        chunks.append(result)
    if diplomas:
        chunks.append(diplomas)
    if awards:
        chunks.append(awards)

    if is_stud and u.mentor:
        chunks.append(f'(наставник {u.mentor.fio})')

    return ' '.join(chunks).strip()


def _group_rows(rows):
    grouped = {}
    order = []

    for row in rows:
        key = row['event_key']
        if key not in grouped:
            grouped[key] = {
                'event_name': row['event_name'],
                'event_level': row['event_level'],
                'event_date': row['event_date'],
                'result_parts': [],
            }
            order.append(key)
        grouped[key]['result_parts'].append(row['result_col'])

    result = []
    for key in order:
        item = grouped[key]
        parts = [p for p in item['result_parts'] if p]
        result.append({
            'event_name': item['event_name'],
            'event_level': item['event_level'],
            'event_date': item['event_date'],
            'result_col': '\n'.join(parts),
        })
    return result


def _fetch_records(year1, year2, report_type, sort):
    query = Ucastie.query.filter_by(year1=year1, year2=year2)

    if sort == 'id':
        query = query.order_by(Ucastie.id)
    elif sort == 'level':
        query = query.join(Ucastie.meropriyatie).join(
            Meropriyatie.uroven
        ).order_by(Uroven.uroven_name, Ucastie.id)
    else:
        query = query.join(Ucastie.meropriyatie).order_by(
            Meropriyatie.date, Ucastie.id
        )

    ped_rows = []
    stud_rows = []

    for u in query.all():
        is_stud = _is_student(u)
        row = {
            'event_key': u.id_meropriyatie,
            'event_name': u.meropriyatie.name,
            'event_level': u.meropriyatie.uroven.uroven_name,
            'event_date': u.meropriyatie.date,
            'result_col': _format_result_column(u, is_stud),
        }
        if is_stud:
            stud_rows.append(row)
        else:
            ped_rows.append(row)

    if report_type == 'ped':
        stud_rows = []
    elif report_type == 'stud':
        ped_rows = []

    return _group_rows(ped_rows), _group_rows(stud_rows)


def _set_cell_text(cell, text, bold=False):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold


def _fill_table(table, rows):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)

    for idx, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        _set_cell_text(cells[0], idx)
        _set_cell_text(cells[1], row['event_name'])
        _set_cell_text(cells[2], row['event_level'])
        _set_cell_text(cells[3], row['event_date'])
        _set_cell_text(cells[4], row['result_col'])


def _replace_title_paragraph(doc, old_fragment, new_text):
    for paragraph in doc.paragraphs:
        if old_fragment in paragraph.text:
            paragraph.text = new_text
            return True
    return False


def _generate_report(year1, year2, report_type, sort):
    from docx import Document

    period = f'{year1}-{year2}'
    ped_rows, stud_rows = _fetch_records(year1, year2, report_type, sort)

    if not ped_rows and not stud_rows:
        return None

    doc = Document(TEMPLATE_PATH)

    _replace_title_paragraph(
        doc,
        'Участие педагогических работников',
        PED_TITLE.format(period=period),
    )
    _replace_title_paragraph(
        doc,
        'Участие студентов',
        STUD_TITLE.format(period=period),
    )

    if len(doc.tables) >= 1 and ped_rows:
        _fill_table(doc.tables[0], ped_rows)
    elif len(doc.tables) >= 1:
        _fill_table(doc.tables[0], [])

    if len(doc.tables) >= 2 and stud_rows:
        _fill_table(doc.tables[1], stud_rows)
    elif len(doc.tables) >= 2:
        _fill_table(doc.tables[1], [])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_report_from_scratch(year1, year2, report_type, sort):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    period = f'{year1}-{year2}'
    ped_rows, stud_rows = _fetch_records(year1, year2, report_type, sort)

    if not ped_rows and not stud_rows:
        return None

    doc = Document()

    header = doc.add_paragraph('АПОУ ВО «ВКСиИТ»')
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].bold = True
    header.runs[0].font.size = Pt(14)

    def add_section(title, rows):
        if not rows:
            return

        p = doc.add_paragraph(title)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(12)

        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, name in enumerate(TABLE_HEADERS):
            _set_cell_text(hdr[i], name, bold=True)

        _fill_table(table, rows)
        doc.add_paragraph('')

    if ped_rows:
        add_section(PED_TITLE.format(period=period), ped_rows)
    if stud_rows:
        add_section(STUD_TITLE.format(period=period), stud_rows)

    for section in doc.sections:
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


@report_bp.route('/download')
def download_report():
    period = request.args.get('year', '').strip()
    report_type = request.args.get('type', 'all').strip()
    sort = request.args.get('sort', 'id').strip()

    if report_type not in ('all', 'ped', 'stud'):
        return {'error': 'Неверный тип отчета'}, 400
    if sort not in ('id', 'date', 'level'):
        return {'error': 'Неверный тип сортировки'}, 400

    try:
        year1, year2 = _parse_period(period)
    except ValueError:
        return {'error': 'Период должен быть в формате XXXX-XXXX'}, 400

    try:
        if report_type == 'all' and os.path.exists(TEMPLATE_PATH):
            buffer = _generate_report(year1, year2, report_type, sort)
        else:
            buffer = _generate_report_from_scratch(year1, year2, report_type, sort)
    except ImportError:
        return {
            'error': 'Установите python-docx: pip install python-docx'
        }, 500

    if buffer is None:
        return {'error': 'Нет данных за выбранный период'}, 404

    filename = f'otchet_{period}_{report_type}.docx'
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            'application/vnd.openxmlformats-officedocument'
            '.wordprocessingml.document'
        ),
    )
